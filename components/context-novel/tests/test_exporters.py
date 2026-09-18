from __future__ import annotations

from zipfile import ZipFile

from docx import Document
from ielts_novel.exporters.docx_exporter import export_chapter_docx, export_volume_docx
from ielts_novel.exporters.html_exporter import export_chapter_html, export_index_html
from ielts_novel.exporters.xlsx_exporter import export_glossary_xlsx
from ielts_novel.models import (
    ConvertedChapter,
    ConvertedParagraph,
    InsertedTerm,
    VocabularyItem,
)
from openpyxl import load_workbook


def _chapter():
    term = InsertedTerm(word="conceal", lemma="conceal", meaning="掩饰", part_of_speech="verb", cefr="B2", phonetic="/kənˈsiːl/", collocation="conceal evidence", example_sentence="He tried to conceal the letter.")
    return ConvertedChapter(chapter_id=1, chapter_title="第一章 山边小村", paragraphs=[ConvertedParagraph(id="1-001", converted_text="她试图 conceal（掩饰）自己的紧张。", inserted_terms=[term])])


def test_docx_has_heading_bold_term_pink_shading_table_and_page_break(tmp_path):
    path = export_chapter_docx(_chapter(), tmp_path / "chapter.docx", cumulative={"conceal": 1})
    doc = Document(path)
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert any(run.text == "conceal" and run.bold for p in doc.paragraphs for run in p.runs)
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "FCE4EC" in xml
    assert "本章核心词汇" in xml
    assert "w:type=\"page\"" in xml


def test_html_has_click_metadata_and_meaning_toggle(tmp_path):
    path = export_chapter_html(_chapter(), tmp_path / "chapter.html", cumulative={"conceal": 1})
    html = path.read_text(encoding="utf-8")
    assert 'class="vocab"' in html
    assert 'data-phonetic="/kənˈsiːl/"' in html
    assert "toggle-meanings" in html
    assert 'class="meaning"' in html
    assert ".hide-meanings .meaning{display:none}" in html
    assert "conceal evidence" in html


def test_xlsx_has_required_columns_filter_and_frozen_header(tmp_path):
    item = VocabularyItem(word="conceal", lemma="conceal", meaning="掩饰", part_of_speech="verb", cefr="B2", occurrence_count=2)
    path = export_glossary_xlsx([item], tmp_path / "glossary.xlsx")
    sheet = load_workbook(path).active
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet["A1"].value == "单词或短语"
    assert sheet["J2"].value == 2


def test_volume_has_toc_field_and_chapter_headings(tmp_path):
    path = export_volume_docx([_chapter(), _chapter().model_copy(update={"chapter_id": 2, "chapter_title": "第二章"})], tmp_path / "volume.docx")
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "TOC \\o" in xml
    assert "第一章 山边小村" in xml and "第二章" in xml


def test_html_index_links_chapters(tmp_path):
    path = export_index_html([(1, "第一章 山边小村")], tmp_path / "index.html")
    html = path.read_text(encoding="utf-8")
    assert "第0001章.html" in html
    assert "第一章 山边小村" in html
