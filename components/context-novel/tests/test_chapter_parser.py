from __future__ import annotations

import json

import pytest
from docx import Document
from ebooklib import epub

from ielts_novel.processors.chapter_parser import ChapterDetectionError, parse_novel, write_detection_report


@pytest.mark.parametrize(
    "title",
    ["第1章 初见", "第一章 初见", "第001章 初见", "Chapter 1 Arrival", "卷一 风起", "番外", "序章", "终章"],
)
def test_txt_recognizes_supported_chapter_titles(tmp_path, title):
    path = tmp_path / "book.txt"
    path.write_text(f"{title}\n第一段。\n第二段。\n第2章 后续\n第三段。", encoding="utf-8")

    result = parse_novel(path)

    assert result.chapters[0].chapter_title == title
    assert [p.text for p in result.chapters[0].paragraphs] == ["第一段。", "第二段。"]
    assert result.confident


def test_txt_falls_back_to_gb18030(tmp_path):
    path = tmp_path / "book.txt"
    path.write_bytes("第一章 开始\n中文正文。\n第二章 继续\n后文。".encode("gb18030"))

    result = parse_novel(path)

    assert result.encoding == "gb18030"
    assert len(result.chapters) == 2


def test_docx_reads_paragraphs(tmp_path):
    path = tmp_path / "book.docx"
    document = Document()
    for text in ["第一章 开始", "段落甲。", "段落乙。", "第二章 后续", "段落丙。"]:
        document.add_paragraph(text)
    document.save(path)

    result = parse_novel(path)

    assert len(result.chapters) == 2
    assert result.chapters[0].paragraphs[1].id == "1-002"


def test_epub_reads_spine_order(tmp_path):
    path = tmp_path / "book.epub"
    book = epub.EpubBook()
    book.set_identifier("fixture")
    book.set_title("fixture")
    chapter = epub.EpubHtml(title="内容", file_name="chapter.xhtml")
    chapter.content = "<h1>第一章 开始</h1><p>段落甲。</p><h1>第二章 后续</h1><p>段落乙。</p>"
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.toc = (chapter,)
    book.spine = [chapter]
    epub.write_epub(path, book)

    result = parse_novel(path)

    assert [item.chapter_title for item in result.chapters] == ["第一章 开始", "第二章 后续"]


def test_epub_ignores_empty_toc_headings_before_real_chapters(tmp_path):
    path = tmp_path / "book-with-inline-toc.epub"
    book = epub.EpubBook()
    book.set_identifier("toc-fixture")
    book.set_title("toc-fixture")
    chapter = epub.EpubHtml(title="内容", file_name="chapter.xhtml")
    chapter.content = (
        "<h1>第一章 开始</h1><h1>第二章 后续</h1>"
        "<h1>第一章 开始</h1><p>正文甲。</p>"
        "<h1>第二章 后续</h1><p>正文乙。</p>"
    )
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.toc = (chapter,)
    book.spine = [chapter]
    epub.write_epub(path, book)

    result = parse_novel(path)

    assert result.confident
    assert len(result.chapters) == 2
    assert result.discarded_empty_chapters == 2
    assert [item.chapter_id for item in result.chapters] == [1, 2]


def test_low_confidence_does_not_return_chapters_and_writes_report(tmp_path):
    path = tmp_path / "broken.txt"
    path.write_text("这是没有章节标题的大段正文。" * 200, encoding="utf-8")
    report_path = tmp_path / "reports" / "chapter_detection.json"

    with pytest.raises(ChapterDetectionError) as caught:
        parse_novel(path)
    write_detection_report(caught.value.result, report_path)

    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["confident"] is False
    assert report["detected_chapters"] == 0
    assert report["source_hash"]
