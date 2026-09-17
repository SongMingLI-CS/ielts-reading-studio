from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from ebooklib import ITEM_DOCUMENT, epub

from ielts_novel.models import Chapter, Paragraph


_CN_NUMBER = "0-9〇零一二三四五六七八九十百千万两"
CHAPTER_TITLE_RE = re.compile(
    rf"^(?:第[{_CN_NUMBER}]+章(?:\s+|$)|Chapter\s+\d+(?:\s+|$)|卷[{_CN_NUMBER}]+(?:\s+|$)|番外(?:\s*[:：]?\s*.*)?$|序章(?:\s+.*)?$|终章(?:\s+.*)?$)",
    re.IGNORECASE,
)


@dataclass(slots=True)
class ParseResult:
    source_path: str
    source_hash: str
    format: str
    encoding: str | None
    chapters: list[Chapter] = field(default_factory=list)
    confident: bool = False
    warnings: list[str] = field(default_factory=list)
    preamble_paragraphs: int = 0

    def report_dict(self) -> dict:
        return {
            "source_path": self.source_path,
            "source_hash": self.source_hash,
            "format": self.format,
            "encoding": self.encoding,
            "confident": self.confident,
            "detected_chapters": len(self.chapters),
            "preamble_paragraphs": self.preamble_paragraphs,
            "warnings": self.warnings,
            "chapter_titles": [chapter.chapter_title for chapter in self.chapters],
        }


class ChapterDetectionError(ValueError):
    def __init__(self, message: str, result: ParseResult):
        super().__init__(message)
        self.result = result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_txt(path: Path) -> tuple[list[str], str]:
    payload = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return payload.decode(encoding).splitlines(), encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeError(f"无法以 UTF-8 或 GB18030 解码：{path.name}")


def _read_docx(path: Path) -> list[str]:
    return [paragraph.text for paragraph in Document(path).paragraphs]


def _read_epub(path: Path) -> list[str]:
    book = epub.read_epub(str(path), options={"ignore_ncx": True})
    by_id = {item.id: item for item in book.get_items_of_type(ITEM_DOCUMENT)}
    lines: list[str] = []
    for item_id, _linear in book.spine:
        item = by_id.get(item_id)
        if item is None:
            continue
        soup = BeautifulSoup(item.get_content(), "html.parser")
        lines.extend(element.get_text(" ", strip=True) for element in soup.find_all(["h1", "h2", "h3", "p"]))
    return lines


def _split_chapters(lines: list[str], path: Path, source_hash: str, file_format: str, encoding: str | None) -> ParseResult:
    chapters: list[Chapter] = []
    current_title: str | None = None
    current_lines: list[str] = []
    preamble = 0

    def flush() -> None:
        nonlocal current_lines
        if current_title is None:
            return
        chapter_id = len(chapters) + 1
        paragraphs = [
            Paragraph(id=f"{chapter_id}-{index:03d}", text=text)
            for index, text in enumerate((line.strip() for line in current_lines if line.strip()), start=1)
        ]
        chapters.append(
            Chapter(
                chapter_id=chapter_id,
                chapter_title=current_title,
                paragraphs=paragraphs,
                source_path=path,
                source_hash=source_hash,
            )
        )
        current_lines = []

    for raw_line in lines:
        line = raw_line.strip().replace("\u3000", " ").strip()
        if CHAPTER_TITLE_RE.match(line):
            flush()
            current_title = line
        elif current_title is None:
            if line:
                preamble += 1
        else:
            current_lines.append(line)
    flush()

    warnings: list[str] = []
    if len(chapters) < 2:
        warnings.append("未识别到至少两个可靠章节边界")
    if chapters and sum(not chapter.paragraphs for chapter in chapters) / len(chapters) > 0.1:
        warnings.append("空章节比例超过 10%")
    confident = len(chapters) >= 2 and not warnings
    return ParseResult(str(path), source_hash, file_format, encoding, chapters, confident, warnings, preamble)


def parse_novel(path: str | Path) -> ParseResult:
    source = Path(path).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    suffix = source.suffix.lower()
    source_hash = _sha256(source)
    if suffix == ".txt":
        lines, encoding = _read_txt(source)
        file_format = "txt"
    elif suffix == ".docx":
        lines, encoding, file_format = _read_docx(source), None, "docx"
    elif suffix == ".epub":
        lines, encoding, file_format = _read_epub(source), None, "epub"
    else:
        raise ValueError(f"不支持的输入格式：{suffix}")
    result = _split_chapters(lines, source, source_hash, file_format, encoding)
    if not result.confident:
        raise ChapterDetectionError("章节识别置信度不足，已停止处理", result)
    return result


def write_detection_report(result: ParseResult, path: str | Path) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".tmp")
    temporary.write_text(json.dumps(result.report_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(target)

