from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ielts_novel.models import VocabularyItem
from ielts_novel.storage.atomic import atomic_save_workbook

HEADERS = ["单词或短语", "lemma", "音标", "词性", "中文释义", "CEFR等级", "IELTS类别", "常用搭配", "原创例句", "本书累计出现次数"]


def export_glossary_xlsx(items: list[VocabularyItem], path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "全局词库"
    sheet.append(HEADERS)
    for item in items:
        sheet.append([item.word, item.lemma, item.phonetic, item.part_of_speech, item.meaning, item.cefr, item.category, item.collocation, item.example_sentence, item.occurrence_count])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [20, 18, 16, 16, 20, 12, 18, 28, 42, 18]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    atomic_save_workbook(workbook, target)
    return target

