from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ielts_novel.models import ConvertedChapter, InsertedTerm
from ielts_novel.storage.atomic import atomic_save_document

PINK = "FCE4EC"


def _set_font(run, latin="Aptos", east_asia="Microsoft YaHei", size=11.5):
    run.font.name = latin
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def _shade(run, fill=PINK):
    props = run._element.get_or_add_rPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def _add_mixed_runs(paragraph, text: str, terms: list[InsertedTerm]):
    by_word = {term.word: term for term in terms}
    if not by_word:
        run = paragraph.add_run(text)
        _set_font(run)
        return
    pattern = re.compile("(" + "|".join(sorted((re.escape(word) for word in by_word), key=len, reverse=True)) + r")(（[^）]+）)?")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            _set_font(run)
        english = paragraph.add_run(match.group(1))
        english.bold = True
        _set_font(english)
        _shade(english)
        if match.group(2):
            meaning = paragraph.add_run(match.group(2))
            _set_font(meaning)
            _shade(meaning)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_font(run)


def export_chapter_docx(chapter: ConvertedChapter, path: str | Path, *, cumulative: dict[str, int] | None = None) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    heading = document.add_heading(chapter.chapter_title, level=1)
    for run in heading.runs:
        _set_font(run, size=18)
        run.font.color.rgb = RGBColor(0, 0, 0)
    unique: dict[str, InsertedTerm] = {}
    for item in chapter.paragraphs:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.55
        _add_mixed_runs(paragraph, item.converted_text, item.inserted_terms)
        for term in item.inserted_terms:
            unique.setdefault(term.lemma, term)
    document.add_heading("本章核心词汇", level=2)
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["单词或短语", "音标", "词性", "文中释义", "CEFR", "常用搭配", "原创例句", "累计次数"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for term in list(unique.values())[:40]:
        row = table.add_row().cells
        values = [term.word, term.phonetic, term.part_of_speech, term.meaning, term.cefr, term.collocation, term.example_sentence, str((cumulative or {}).get(term.lemma, 1))]
        for cell, value in zip(row, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    atomic_save_document(document, target)
    return target


def export_volume_docx(chapters: list[ConvertedChapter], path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    title = document.add_paragraph("分卷目录", style="Title")
    for run in title.runs:
        _set_font(run, size=20)
        run.font.color.rgb = RGBColor(0, 0, 0)
    toc_paragraph = document.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-1" \\h \\z \\u')
    toc_paragraph._p.append(field)
    document.add_page_break()
    for chapter_index, chapter in enumerate(chapters):
        heading = document.add_heading(chapter.chapter_title, level=1)
        for run in heading.runs:
            _set_font(run, size=18)
            run.font.color.rgb = RGBColor(0, 0, 0)
        for item in chapter.paragraphs:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.55
            _add_mixed_runs(paragraph, item.converted_text, item.inserted_terms)
        if chapter_index < len(chapters) - 1:
            document.add_page_break()
    atomic_save_document(document, target)
    return target
