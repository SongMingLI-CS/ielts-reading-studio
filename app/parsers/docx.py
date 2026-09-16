from __future__ import annotations

from pathlib import Path

from docx import Document

from .base import ParseResult
from .chapter_detection import (
    ChapterDraft,
    DetectedHeading,
    HeadingKind,
    detect_heading,
    finalize_chapters,
    resolve_units,
)

HEADING_STYLE_PREFIXES = ("heading", "标题")


def parse_docx(path: Path, source_hash: str) -> ParseResult:
    """Parse a Word source using heading styles, falling back to title patterns."""
    document = Document(str(path))
    entries = _entries(document)
    styled_structure = any(is_heading for _, is_heading in entries)
    units = _collect(entries, styled_structure)
    chapters, preamble_chars = resolve_units(units)
    finalized = finalize_chapters(
        source_hash, chapters, explicit_structure=styled_structure, preamble_chars=preamble_chars,
    )
    return ParseResult(
        format="docx",
        encoding=None,
        source_hash=source_hash,
        chapters=finalized.chapters,
        confidence=finalized.confidence,
        diagnostics=finalized.diagnostics,
        candidate_chapters=finalized.candidates,
    )


def _entries(document) -> list[tuple[str, bool]]:
    """Return non-empty body paragraphs paired with their heading-style flag."""
    entries: list[tuple[str, bool]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        entries.append((text, _is_heading_style(paragraph)))
    return entries


def _is_heading_style(paragraph) -> bool:
    name = getattr(getattr(paragraph, "style", None), "name", "") or ""
    return name.strip().lower().startswith(HEADING_STYLE_PREFIXES)


def _collect(entries: list[tuple[str, bool]], styled_structure: bool) -> list[ChapterDraft]:
    units: list[ChapterDraft] = []
    current = ChapterDraft(heading=None, offsets={"start": 0, "end": 0})
    current_start = 0
    offset = 0
    for text, is_heading in entries:
        is_boundary = is_heading if styled_structure else detect_heading(text) is not None
        if is_boundary:
            current.offsets = {"start": current_start, "end": offset}
            units.append(current)
            current = ChapterDraft(heading=_heading(text))
            current_start = offset
        else:
            current.paragraphs.append(text)
        offset += len(text) + 1
    current.offsets = {"start": current_start, "end": offset}
    units.append(current)
    return units


def _heading(text: str) -> DetectedHeading:
    normalized = text.strip()
    detected = detect_heading(normalized)
    if detected is not None:
        return detected
    return DetectedHeading(HeadingKind.CHAPTER, None, normalized)
