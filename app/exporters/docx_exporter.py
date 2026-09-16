from __future__ import annotations

import os
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_BREAK

from app.models import ReadingPackage

from .json_exporter import _require_validated


def export_docx(package: ReadingPackage, path: str | Path) -> Path:
    _require_validated(package)
    document = Document()
    _add_package(document, package)
    return _save_atomic(document, Path(path))


def export_workbooks(
    packages: list[ReadingPackage],
    directory: str | Path,
    *,
    size: int = 20,
) -> list[Path]:
    if not 20 <= size <= 50:
        raise ValueError("Workbook size must be between 20 and 50 packages")
    for package in packages:
        _require_validated(package)
    output_dir = Path(directory)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for start in range(0, len(packages), size):
        document = Document()
        batch = packages[start : start + size]
        for index, package in enumerate(batch):
            if index:
                document.add_page_break()
            _add_package(document, package)
        volume = start // size + 1
        paths.append(_save_atomic(document, output_dir / f"workbook-{volume:03d}.docx"))
    return paths


def _add_package(document: DocumentType, package: ReadingPackage) -> None:
    passage = package.passage
    document.add_heading("IELTS Academic Reading Practice", level=1)
    document.add_paragraph(f"Difficulty: {passage.difficulty.value.title()}")
    document.add_heading(passage.title, level=2)
    for paragraph in passage.paragraphs:
        body = document.add_paragraph()
        body.add_run(f"{paragraph.label}  ").bold = True
        body.add_run(paragraph.text)

    document.add_heading("Questions", level=2)
    for group in package.question_groups:
        document.add_heading(group.type.value.replace("_", " ").title(), level=3)
        instructions = group.instructions
        if group.word_limit:
            instructions += f" (NO MORE THAN {group.word_limit} WORDS)"
        document.add_paragraph(instructions)
        if group.options:
            document.add_paragraph("Options: " + " | ".join(group.options))
        for question in group.questions:
            document.add_paragraph(f"{question.number}. {question.prompt}\nAnswer: ____________________")

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Answer Key and Analysis", level=2)
    for group in package.question_groups:
        for question in group.questions:
            document.add_heading(f"{question.number}. {question.answer}", level=3)
            document.add_paragraph(
                f"Evidence ({question.evidence_paragraph}): {question.evidence_quote}"
            )
            document.add_paragraph(question.chinese_explanation)
            for option, explanation in question.distractor_explanations.items():
                document.add_paragraph(f"{option}: {explanation}")

    document.add_heading("Vocabulary", level=2)
    table = document.add_table(rows=1, cols=6)
    headings = ["Word", "Pronunciation", "Part of speech", "中文", "Collocations", "Example"]
    for cell, heading in zip(table.rows[0].cells, headings):
        cell.text = heading
    for entry in passage.vocabulary:
        cells = table.add_row().cells
        values = [
            entry.word,
            entry.pronunciation or "",
            entry.part_of_speech or "",
            entry.chinese_meaning or "",
            ", ".join(entry.collocations),
            entry.example or "",
        ]
        for cell, value in zip(cells, values):
            cell.text = value


def _save_atomic(document: DocumentType, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.parent / f"{destination.stem}.{uuid4().hex}.tmp.docx"
    try:
        document.save(temporary)
        os.replace(temporary, destination)
    finally:
        if temporary.exists():
            temporary.unlink()
    return destination

