from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from app.parsers.base import SourceNotFoundError, UnsupportedFormatError
from app.parsers.registry import parse_source

FIXTURES = Path(__file__).parent.parent / "fixtures"


def test_txt_falls_back_to_gb18030(tmp_path):
    path = tmp_path / "book.txt"
    path.write_bytes("第一章 开始\n正文。\n第二章 后续\n后文。".encode("gb18030"))

    result = parse_source(path)

    assert result.encoding == "gb18030"
    assert result.format == "txt"
    assert result.confident
    assert [c.chapter_title for c in result.chapters] == ["第一章 开始", "第二章 后续"]
    assert result.chapters[0].ordinal == 1
    assert result.chapters[0].paragraphs[0].text == "正文。"
    assert result.chapters[0].character_count == 3
    assert result.diagnostics == []


def test_txt_source_hash_matches_file_bytes(tmp_path):
    path = tmp_path / "book.txt"
    payload = "第一章 开始\n正文。\n第二章 后续\n后文。".encode()
    path.write_bytes(payload)

    result = parse_source(path)

    assert result.source_hash == hashlib.sha256(payload).hexdigest()


def test_txt_detects_utf8_bom_and_reports_offsets(tmp_path):
    path = tmp_path / "book.txt"
    path.write_bytes("\ufeff第一章 开始\n正文。\n第二章 后续\n后文。".encode("utf-8"))

    result = parse_source(path)

    assert result.encoding == "utf-8-sig"
    assert [c.chapter_title for c in result.chapters] == ["第一章 开始", "第二章 后续"]
    first, second = result.chapters
    assert first.source_offsets["start"] < second.source_offsets["start"]
    assert first.source_offsets["end"] == second.source_offsets["start"]


def test_txt_tracks_volume_titles(tmp_path):
    path = tmp_path / "book.txt"
    path.write_text(
        "卷一 起风\n第一章 开始\n正文。\n第二章 后续\n后文。\n卷二 落雨\n第三章 收束\n结尾。",
        encoding="utf-8",
    )

    result = parse_source(path)

    assert [c.chapter_title for c in result.chapters] == ["第一章 开始", "第二章 后续", "第三章 收束"]
    assert [c.volume_title for c in result.chapters] == ["卷一 起风", "卷一 起风", "卷二 落雨"]
    assert [c.ordinal for c in result.chapters] == [1, 2, 3]


def test_txt_identity_and_offsets_are_stable_between_parses(tmp_path):
    path = tmp_path / "book.txt"
    path.write_text("第一章 开始\n正文。\n第二章 后续\n后文。", encoding="utf-8")

    first = parse_source(path)
    second = parse_source(path)

    assert [c.id for c in first.chapters] == [c.id for c in second.chapters]
    assert [c.source_offsets for c in first.chapters] == [c.source_offsets for c in second.chapters]


def test_markdown_uses_headings_as_boundaries(tmp_path):
    path = tmp_path / "book.md"
    path.write_text("# 第一章\n甲。\n# 第二章\n乙。", encoding="utf-8")

    result = parse_source(path)

    assert result.format == "md"
    assert len(result.chapters) == 2
    assert [c.chapter_title for c in result.chapters] == ["第一章", "第二章"]


def test_markdown_fixture_parses_three_chapters():
    result = parse_source(FIXTURES / "sample.md")

    assert result.confident
    assert [c.chapter_title for c in result.chapters] == [
        "第一章 起步", "第二章 深入", "第三章 扩展",
    ]
    assert result.chapters[0].paragraphs[0].text == "水源的分布决定了聚落的形态。"


def test_docx_reads_non_empty_body_paragraphs_and_heading_styles(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("前言部分。")
    document.add_paragraph("")
    document.add_heading("第一章 开端", level=1)
    document.add_paragraph("甲。")
    document.add_heading("第二章 结束", level=1)
    document.add_paragraph("乙。")
    path = tmp_path / "book.docx"
    document.save(path)

    result = parse_source(path)

    assert result.format == "docx"
    assert result.encoding is None
    assert result.confident
    assert [c.chapter_title for c in result.chapters] == ["第一章 开端", "第二章 结束"]
    assert result.chapters[0].paragraphs[0].text == "甲。"


def test_docx_without_heading_styles_uses_title_patterns(tmp_path):
    from docx import Document

    document = Document()
    for text in ("第一章 开端", "甲。", "乙。", "第二章 结束", "丙。"):
        document.add_paragraph(text)
    path = tmp_path / "plain.docx"
    document.save(path)

    result = parse_source(path)

    assert [c.chapter_title for c in result.chapters] == ["第一章 开端", "第二章 结束"]
    assert [p.text for p in result.chapters[0].paragraphs] == ["甲。", "乙。"]


def test_epub_follows_spine_order(tmp_path):
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("id-1")
    book.set_title("Book")
    book.set_language("zh")
    first = epub.EpubHtml(title="第一章", file_name="c1.xhtml", lang="zh")
    first.content = "<html><body><h1>第一章 开端</h1><p>甲。</p></body></html>"
    second = epub.EpubHtml(title="第二章", file_name="c2.xhtml", lang="zh")
    second.content = "<html><body><h1>第二章 结束</h1><p>乙。</p></body></html>"
    book.add_item(first)
    book.add_item(second)
    book.toc = (first, second)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", first, second]
    path = tmp_path / "book.epub"
    epub.write_epub(str(path), book)

    result = parse_source(path)

    assert result.format == "epub"
    assert result.confident
    assert [c.chapter_title for c in result.chapters] == ["第一章 开端", "第二章 结束"]
    assert result.chapters[1].paragraphs[0].text == "乙。"


def test_unstructured_large_text_returns_diagnostics_not_one_giant_chapter(tmp_path):
    path = tmp_path / "broken.txt"
    path.write_text("没有章节边界的正文。" * 1000, encoding="utf-8")

    result = parse_source(path)

    assert not result.confident
    assert result.chapters == []
    assert "no_reliable_boundaries" in result.diagnostics


def test_txt_with_a_single_heading_is_not_confident(tmp_path):
    path = tmp_path / "single.txt"
    path.write_text("第一章 唯一\n" + "正文。" * 500, encoding="utf-8")

    result = parse_source(path)

    assert not result.confident
    assert result.chapters == []
    assert result.candidate_chapters
    assert "no_reliable_boundaries" in result.diagnostics


def test_unsupported_extension_is_rejected(tmp_path):
    path = tmp_path / "book.pdf"
    path.write_text("irrelevant", encoding="utf-8")

    with pytest.raises(UnsupportedFormatError, match=r"\.pdf"):
        parse_source(path)


def test_missing_source_is_rejected(tmp_path):
    with pytest.raises(SourceNotFoundError, match="not found"):
        parse_source(tmp_path / "missing.txt")
