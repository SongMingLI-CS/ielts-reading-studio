from copy import deepcopy

import pytest
from docx import Document

from app.exporters.docx_exporter import export_docx, export_workbooks


def test_docx_contains_passage_questions_answer_key_and_vocabulary(valid_package, tmp_path):
    path = export_docx(valid_package, tmp_path / "practice.docx")
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert valid_package.passage.title in text
    assert "Questions" in text
    assert "Answer Key and Analysis" in text
    assert valid_package.question_groups[0].questions[0].chinese_explanation in text
    assert document.tables
    assert document.tables[-1].cell(1, 0).text == "supply"


def test_export_workbooks_splits_51_packages_into_bounded_volumes(valid_package, tmp_path):
    packages = []
    for index in range(51):
        package = deepcopy(valid_package)
        package.unit.id = f"unit-{index + 1}"
        package.passage.title = f"Passage {index + 1}"
        packages.append(package)

    paths = export_workbooks(packages, tmp_path, size=20)

    assert len(paths) == 3
    texts = ["\n".join(p.text for p in Document(path).paragraphs) for path in paths]
    assert sum(text.count("IELTS Academic Reading Practice") for text in texts) == 51
    assert "Passage 20" in texts[0] and "Passage 21" not in texts[0]
    assert "Passage 40" in texts[1] and "Passage 41" not in texts[1]
    assert "Passage 51" in texts[2]


@pytest.mark.parametrize("size", [0, 19, 51, 100])
def test_export_workbooks_rejects_out_of_bounds_size(valid_package, tmp_path, size):
    with pytest.raises(ValueError, match="20.*50"):
        export_workbooks([valid_package], tmp_path, size=size)
